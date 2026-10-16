"""
Flatten results into rows and write them to Excel, one sheet per
document type. Data-element orders give one row per employee.
"""

from __future__ import annotations

import time
from contextlib import contextmanager
from datetime import datetime
from pathlib import Path

try:
    import msvcrt
except ImportError:  # not Windows
    import fcntl
    msvcrt = None

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from go_extractor.data_elements import ORDER_TYPES, elements_for
from go_extractor.models import DocumentResult, EmployeeElements

# (column, width) per document type.
ADVOCATE_COLUMNS = [
    ("Document", 28),
    ("File No.", 18),
    ("Govt. Order No.", 18),
    ("Govt. Order Date", 14),
    ("Case No.", 22),
    ("Litigant (Petitioner)", 30),
    ("Respondent", 30),
    ("Respondents Represented", 50),
    ("Government Advocate", 40),
    ("Litigation Conducting Officer", 50),
    ("Swearing Officer", 22),
    ("LCO Phone", 16),
    ("LCO Email", 30),
    ("Issuing Authority", 24),
    ("Issuing Authority Designation", 50),
    ("Status", 14),
    ("Processed At", 20),
]

SUSPENSION_COLUMNS = [
    ("Document", 34),
    ("Order Type", 22),
    ("Order No.", 30),
    ("Order Date", 12),
    ("Employee Name", 26),
    ("Designation", 30),
    ("Office", 40),
    ("Reason", 50),
    ("Rule Invoked", 30),
    ("Case Reference", 30),
    ("Effective From", 18),
    ("Issuing Authority", 24),
    ("Issuing Authority Designation", 40),
    ("Status", 14),
    ("Processed At", 20),
    # As written in the order, for checking against the scan. Kept last so
    # sheets created before these columns existed can be extended.
    ("Order No. (Kannada)", 30),
    ("Employee Name (Kannada)", 26),
    ("Designation (Kannada)", 30),
    ("Office (Kannada)", 40),
    ("Rule Invoked (Kannada)", 30),
    ("Case Reference (Kannada)", 30),
    ("Issuing Authority (Kannada)", 24),
    ("Issuing Authority Designation (Kannada)", 40),
]


def _element_columns(order_type: str) -> list[tuple[str, int]]:
    """
    Document, employee "1 of 3", a warning when the subject line reads as
    another order type than the file name, the data elements, Status, then
    the as-written values.
    """
    elements = elements_for(order_type)

    return [
        ("Document", 34),
        ("Employee", 10),
        ("Type Check", 30),
        *((e.label, 14 if e.kind in ("date", "yes_no") else 30)
          for e in elements),
        ("Status", 14),
        ("Processed At", 20),
        *((f"{e.label} (Kannada)", 30)
          for e in elements if e.kind in ("text", "list")),
    ]


SHEETS = {
    "advocate": ("GO Extractions", ADVOCATE_COLUMNS),
    "suspension": ("Suspension Orders", SUSPENSION_COLUMNS),
    **{
        order_type: (f"{order_type}s", _element_columns(order_type))
        for order_type in ORDER_TYPES.values()
    },
}


def to_rows(result: DocumentResult) -> list[dict[str, str]]:
    """One row per document, or per employee for data-element orders."""
    if result.doc_type in ORDER_TYPES.values():
        rows = _element_rows(result)
    elif result.doc_type == "suspension":
        rows = [_suspension_row(result)]
    else:
        rows = [_advocate_row(result)]

    status = f"Failed: {result.error}" if result.error else "Success"
    processed = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    for row in rows:
        row["Status"] = status
        row["Processed At"] = processed

    return rows


def _advocate_row(result: DocumentResult) -> dict[str, str]:
    f = result.fields
    lco = f.litigation_officer

    if lco.is_swearing_officer:
        swearing = "Same as Litigation Conducting Officer"
    else:
        swearing = lco.swearing_officer or ""

    return {
        "Document": result.filename,
        "File No.": f.file_no or "",
        "Govt. Order No.": f.go_number or "",
        "Govt. Order Date": f.go_date or "",
        "Case No.": f.case_number or "",
        "Litigant (Petitioner)": f.petitioner or "",
        "Respondent": f.respondent or "",
        "Respondents Represented": "\n".join(
            f"{i}) {r}" for i, r in enumerate(f.respondents_represented, 1)
        ),
        "Government Advocate": f.government_advocate or "",
        "Litigation Conducting Officer": lco.officer or "",
        "Swearing Officer": swearing,
        "LCO Phone": lco.phone or "",
        "LCO Email": lco.email or "",
        "Issuing Authority": f.issuing_authority.name or "",
        "Issuing Authority Designation": f.issuing_authority.designation or "",
    }


def _suspension_row(result: DocumentResult) -> dict[str, str]:
    f = result.fields

    return {
        "Document": result.filename,
        "Order Type": f.order_type or "",
        "Order No.": f.order_number or "",
        "Order Date": f.order_date or "",
        "Employee Name": f.employee_name or "",
        "Designation": f.employee_designation or "",
        "Office": f.employee_office or "",
        "Reason": f.reason or "",
        "Rule Invoked": f.rule_invoked or "",
        "Case Reference": f.case_reference or "",
        "Effective From": f.effective_from or "",
        "Issuing Authority": f.issuing_authority.name or "",
        "Issuing Authority Designation": f.issuing_authority.designation or "",
        **{
            # getattr: results made before "original" existed (a Streamlit
            # session kept across a code reload) have no such field.
            f"{column} (Kannada)": (getattr(f, "original", None) or {}).get(
                key, ""
            )
            for column, key in _ORIGINAL_COLUMNS
        },
    }


def _element_rows(result: DocumentResult) -> list[dict[str, str]]:
    """The order's values repeated on each employee's row."""
    f = result.fields
    elements = elements_for(result.doc_type)
    # An order where no employee was found still gets one row.
    employees = f.employees or [EmployeeElements()]
    rows = []

    for number, employee in enumerate(employees, start=1):
        values = {**f.values, **employee.values}
        original = {**f.original, **employee.original}

        rows.append({
            "Document": result.filename,
            "Employee": f"{number} of {len(employees)}",
            "Type Check": getattr(f, "type_check", None) or "",
            **{e.label: values.get(e.key) or "" for e in elements},
            **{f"{e.label} (Kannada)": original.get(e.key, "")
               for e in elements if e.kind in ("text", "list")},
        })

    return rows


_ORIGINAL_COLUMNS = [
    ("Order No.", "order_number"),
    ("Employee Name", "employee_name"),
    ("Designation", "employee_designation"),
    ("Office", "employee_office"),
    ("Rule Invoked", "rule_invoked"),
    ("Case Reference", "case_reference"),
    ("Issuing Authority", "issuing_authority"),
    ("Issuing Authority Designation", "issuing_authority_designation"),
]


@contextmanager
def excel_lock(path: Path):
    """One program at a time reads or saves the Excel file.

    Several runs (terminals) can then save to the same file: each one
    waits for the others' saves instead of overwriting their rows. The
    lock is on <file>.lock next to it and is released if a run is killed.
    """
    lock = path.with_name(path.name + ".lock")

    with open(lock, "a+b") as handle:
        while True:
            try:
                handle.seek(0)
                if msvcrt:
                    msvcrt.locking(handle.fileno(), msvcrt.LK_NBLCK, 1)
                else:
                    fcntl.flock(handle, fcntl.LOCK_EX | fcntl.LOCK_NB)
                break
            except OSError:
                time.sleep(0.2)

        try:
            yield
        finally:
            handle.seek(0)
            if msvcrt:
                msvcrt.locking(handle.fileno(), msvcrt.LK_UNLCK, 1)
            else:
                fcntl.flock(handle, fcntl.LOCK_UN)


def append_to_excel(results: list[DocumentResult], path: Path) -> None:
    """Append each document's rows to its type's sheet (created if needed)."""
    with excel_lock(path):
        if path.exists():
            workbook = load_workbook(path)
        else:
            workbook = Workbook()
            workbook.remove(workbook.active)

        for result in results:
            title, columns = SHEETS[result.doc_type]
            sheet = _sheet(workbook, title, columns)

            for row in to_rows(result):
                sheet.append([row[name] for name, _ in columns])

                for cell in sheet[sheet.max_row]:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")

        if not workbook.sheetnames:
            _sheet(workbook, *SHEETS["advocate"])

        workbook.save(path)


def _sheet(workbook: Workbook, title: str, columns: list[tuple[str, int]]):
    if title in workbook.sheetnames:
        sheet = workbook[title]
    else:
        sheet = workbook.create_sheet(title)

    # New sheet, or one written before columns were added at the end.
    for index, (name, width) in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=index)

        if cell.value is None:
            cell.value = name
            cell.font = Font(bold=True)
            sheet.column_dimensions[get_column_letter(index)].width = width

    return sheet


def to_excel_bytes(results: list[DocumentResult]) -> bytes:
    """A fresh workbook with just these results (for downloads)."""
    import tempfile

    with tempfile.TemporaryDirectory() as folder:
        path = Path(folder) / "go_extractions.xlsx"
        append_to_excel(results, path)
        return path.read_bytes()
